"""Document loading: turns uploaded files into LangChain Document objects."""

from pathlib import Path

from docx import Document as DocxDocument
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_core.documents import Document
from pptx import Presentation

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx", ".pptx"}


def _load_docx(file_path: str) -> list[Document]:
    doc = DocxDocument(file_path)
    parts: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    content = "\n\n".join(parts).strip()
    if not content:
        raise ValueError("Word document contains no extractable text.")

    return [
        Document(
            page_content=content,
            metadata={"source": file_path, "file_type": "docx"},
        )
    ]


def _load_pptx(file_path: str) -> list[Document]:
    presentation = Presentation(file_path)
    documents: list[Document] = []

    for index, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = str(shape.text).strip()
                if text:
                    parts.append(text)
        content = "\n".join(parts).strip()
        if content:
            documents.append(
                Document(
                    page_content=content,
                    metadata={
                        "source": file_path,
                        "file_type": "pptx",
                        "page": index,
                        "slide": index,
                    },
                )
            )

    if not documents:
        raise ValueError("PowerPoint file contains no extractable text.")

    return documents


def load_document(file_path: str) -> list[Document]:
    """Load a supported office/text file into LangChain Document objects."""
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix in {".doc", ".ppt"}:
        raise ValueError(
            f"Legacy '{suffix}' files are not supported. "
            "Please convert to .docx or .pptx and upload again."
        )

    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{suffix}'. "
            f"Allowed: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if suffix == ".pdf":
        return PyPDFLoader(str(path)).load()
    if suffix == ".docx":
        return _load_docx(str(path))
    if suffix == ".pptx":
        return _load_pptx(str(path))

    return TextLoader(str(path), encoding="utf-8").load()
